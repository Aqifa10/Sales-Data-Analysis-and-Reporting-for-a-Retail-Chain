
# PHASE 4 — Excel Dashboard & Reports
import pandas as pd
import sqlite3
import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

# ── Paths ─────────────────────────────────────────────────
BASE_DIR = r"C:\Users\KIIT\OneDrive\Desktop\data_analytic\Sales Data Analysis and Reporting for a Retail Chain"
DATA_DIR = os.path.join(BASE_DIR, "data")
IMG_DIR  = os.path.join(BASE_DIR, "images")
RPT_DIR  = os.path.join(BASE_DIR, "reports")
DB_PATH  = os.path.join(DATA_DIR, "retail_analytics.db")
OUT_FILE = os.path.join(RPT_DIR, "Retail_Sales_Report.xlsx")

# ── Load Data ─────────────────────────────────────────────
print("Loading data...")
conn = sqlite3.connect(DB_PATH)
df   = pd.read_sql_query("SELECT * FROM fact_transactions_clean", conn)
rfm  = pd.read_sql_query("SELECT * FROM rfm_segments", conn)
conn.close()

df['trans_date'] = pd.to_datetime(df['trans_date'])
print(f"Loaded {len(df):,} rows")

# ── Summary Tables ────────────────────────────────────────
# Monthly
monthly = df.groupby(['year','month_num','month_name','year_month']).agg(
    transactions     =('tran_amount','count'),
    total_revenue    =('tran_amount','sum'),
    avg_revenue      =('tran_amount','mean'),
    unique_customers =('customer_id','nunique')
).reset_index().sort_values(['year','month_num'])
monthly['avg_revenue'] = monthly['avg_revenue'].round(2)

# Yearly
yearly = df.groupby('year').agg(
    transactions     =('tran_amount','count'),
    total_revenue    =('tran_amount','sum'),
    avg_revenue      =('tran_amount','mean'),
    unique_customers =('customer_id','nunique')
).reset_index()
yearly['avg_revenue'] = yearly['avg_revenue'].round(2)

# Segments
segments = df.groupby('customer_segment').agg(
    customers    =('customer_id','nunique'),
    transactions =('tran_amount','count'),
    total_revenue=('tran_amount','sum'),
    avg_revenue  =('tran_amount','mean')
).reset_index().sort_values('total_revenue', ascending=False)
segments['avg_revenue'] = segments['avg_revenue'].round(2)

# Day of week
dow = df.groupby(['day_of_week','day_name']).agg(
    transactions =('tran_amount','count'),
    total_revenue=('tran_amount','sum'),
    avg_revenue  =('tran_amount','mean')
).reset_index().sort_values('day_of_week')
dow['avg_revenue'] = dow['avg_revenue'].round(2)

# RFM summary
rfm_summary = rfm.groupby('segment').agg(
    customers    =('customer_id','count'),
    avg_recency  =('recency','mean'),
    avg_frequency=('frequency','mean'),
    avg_monetary =('monetary','mean'),
    total_revenue=('monetary','sum')
).reset_index().sort_values('total_revenue', ascending=False).round(1)

# Quarterly
quarterly = df.groupby(['year','quarter']).agg(
    transactions =('tran_amount','count'),
    total_revenue=('tran_amount','sum'),
    avg_revenue  =('tran_amount','mean')
).reset_index()
quarterly['label']       = quarterly['year'].astype(str) + ' Q' + quarterly['quarter'].astype(str)
quarterly['avg_revenue'] = quarterly['avg_revenue'].round(2)

# Top 20 customers
top_custs = df.groupby('customer_id').agg(
    transactions  =('tran_amount','count'),
    total_spent   =('tran_amount','sum'),
    avg_spent     =('tran_amount','mean'),
    first_purchase=('trans_date','min'),
    last_purchase =('trans_date','max')
).reset_index().sort_values('total_spent', ascending=False).head(20)
top_custs['avg_spent'] = top_custs['avg_spent'].round(2)

# Response
resp_summary = df.groupby('responded').agg(
    customers    =('customer_id','nunique'),
    transactions =('tran_amount','count'),
    total_revenue=('tran_amount','sum'),
    avg_revenue  =('tran_amount','mean')
).reset_index()
resp_summary['avg_revenue'] = resp_summary['avg_revenue'].round(2)

print("Summary tables built")

# ── Style Helpers ─────────────────────────────────────────
ORANGE = "E85D04"
AMBER  = "F48C06"
DARK   = "0A0A0F"
CARD   = "1A1A26"
HEADER = "12121A"
WHITE  = "FFFFFF"
LIGHT  = "F0F0F5"
GREEN  = "52B788"
GREY   = "888899"

def fill(c):
    return PatternFill("solid", fgColor=c)

def fnt(bold=False, size=11, color=WHITE):
    return Font(bold=bold, size=size, color=color, name="Calibri")

def aln(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def thin_border():
    s = Side(style="thin", color="2a2a3a")
    return Border(left=s, right=s, top=s, bottom=s)

def dark_bg(ws, rows, cols):
    for r in range(1, rows):
        for c in range(1, cols):
            ws.cell(r, c).fill = fill(DARK)

def write_title(ws, row, col, text, end_col, size=14, color=ORANGE):
    ws.merge_cells(start_row=row, start_column=col,
                   end_row=row, end_column=end_col)
    cell = ws.cell(row, col, text)
    cell.fill      = fill(DARK)
    cell.font      = fnt(bold=True, size=size, color=color)
    cell.alignment = aln(h="left")
    ws.row_dimensions[row].height = 35

def write_headers(ws, row, col, headers, bg=ORANGE):
    for i, h in enumerate(headers, col):
        c = ws.cell(row, i, h)
        c.fill      = fill(bg)
        c.font      = fnt(bold=True, size=10)
        c.alignment = aln()
        c.border    = thin_border()

def write_row(ws, row, col, values, bg=CARD):
    for i, v in enumerate(values, col):
        c = ws.cell(row, i, v)
        c.fill      = fill(bg)
        c.font      = fnt(size=10, color=LIGHT)
        c.alignment = aln()
        c.border    = thin_border()

def set_widths(ws, widths, start=1):
    for i, w in enumerate(widths, start):
        ws.column_dimensions[get_column_letter(i)].width = w

# ── Build Workbook ────────────────────────────────────────
print("Building Excel workbook...")
wb = Workbook()
wb.remove(wb.active)

# ══════════════════════════════════════════════
# SHEET 1 — DASHBOARD
# ══════════════════════════════════════════════
ws = wb.create_sheet("DASHBOARD")
ws.sheet_view.showGridLines = False
dark_bg(ws, 80, 22)

# Title
ws.merge_cells("A1:T1")
ws["A1"] = "RETAIL CHAIN — SALES ANALYTICS DASHBOARD"
ws["A1"].fill      = fill(ORANGE)
ws["A1"].font      = fnt(bold=True, size=20)
ws["A1"].alignment = aln()
ws.row_dimensions[1].height = 45

# Subtitle
ws.merge_cells("A2:T2")
ws["A2"] = "Dataset: 2011–2015   |   125,000 Transactions   |   6,889 Customers   |   $8,123,989 Total Revenue"
ws["A2"].fill      = fill(HEADER)
ws["A2"].font      = fnt(size=10, color=GREY)
ws["A2"].alignment = aln()
ws.row_dimensions[2].height = 22

# KPI Cards row 4-8
ws.row_dimensions[3].height = 10
kpis = [
    ("TOTAL REVENUE",     "$8,123,989",  "A", "D"),
    ("TOTAL ORDERS",      "125,000",     "E", "H"),
    ("AVG TRANSACTION",   "$64.99",      "I", "L"),
    ("UNIQUE CUSTOMERS",  "6,889",       "M", "P"),
    ("RESPONSE RATE",     "9.4%",        "Q", "T"),
]
for label, value, c1, c2 in kpis:
    # Label row
    ws.merge_cells(f"{c1}4:{c2}4")
    lc = ws[f"{c1}4"]
    lc.value     = label
    lc.fill      = fill(HEADER)
    lc.font      = fnt(size=9, color=GREY, bold=True)
    lc.alignment = aln()

    # Value rows
    ws.merge_cells(f"{c1}5:{c2}7")
    vc = ws[f"{c1}5"]
    vc.value     = value
    vc.fill      = fill(CARD)
    vc.font      = fnt(size=18, color=ORANGE, bold=True)
    vc.alignment = aln()

    # Bottom accent
    ws.merge_cells(f"{c1}8:{c2}8")
    bc = ws[f"{c1}8"]
    bc.fill = fill(ORANGE)

for r in [4,8]:
    ws.row_dimensions[r].height = 16
for r in [5,6,7]:
    ws.row_dimensions[r].height = 20

# Insights section
ws.row_dimensions[9].height = 10
ws.merge_cells("A10:T10")
ws["A10"] = "  KEY INSIGHTS"
ws["A10"].fill      = fill(AMBER)
ws["A10"].font      = fnt(bold=True, size=11, color=DARK)
ws["A10"].alignment = aln(h="left")
ws.row_dimensions[10].height = 22

insights = [
    "  ▶   Revenue is consistent across 2011–2015, averaging $1.6M per year with stable growth.",
    "  ▶   Thursday is the highest revenue day — best day to run promotions and campaigns.",
    "  ▶   Platinum customers generate 40% of total revenue — prioritise retention for this group.",
    "  ▶   Only 9.4% of customers responded to campaigns — Champions segment responds 3x more.",
    "  ▶   Transaction amounts cluster between $50–$80 — strong upsell opportunity above $80.",
]
for i, text in enumerate(insights, 11):
    ws.merge_cells(f"A{i}:T{i}")
    c = ws[f"A{i}"]
    c.value     = text
    c.fill      = fill(CARD if i % 2 == 0 else HEADER)
    c.font      = fnt(size=10, color=LIGHT)
    c.alignment = aln(h="left")
    ws.row_dimensions[i].height = 20

# Charts label
ws.row_dimensions[16].height = 10
ws.merge_cells("A17:T17")
ws["A17"] = "  VISUALIZATIONS"
ws["A17"].fill      = fill(HEADER)
ws["A17"].font      = fnt(bold=True, size=11, color=ORANGE)
ws["A17"].alignment = aln(h="left")
ws.row_dimensions[17].height = 22

# Embed charts
embeds = [
    ("01_monthly_revenue.png",  "A18",  480, 250),
    ("02_yearly_revenue.png",   "K18",  480, 250),
    ("05_customer_segments.png","A36",  480, 250),
    ("08_rfm_segments.png",     "K36",  480, 250),
    ("09_cohort_retention.png", "A54",  480, 250),
    ("12_quarterly_growth.png", "K54",  480, 250),
]
for fname, anchor, w, h in embeds:
    fpath = os.path.join(IMG_DIR, fname)
    if os.path.exists(fpath):
        img        = XLImage(fpath)
        img.width  = w
        img.height = h
        ws.add_image(img, anchor)

set_widths(ws, [10]*20)
print("  ✔ DASHBOARD sheet done")

# ══════════════════════════════════════════════
# SHEET 2 — MONTHLY REVENUE
# ══════════════════════════════════════════════
ws2 = wb.create_sheet("Monthly Revenue")
ws2.sheet_view.showGridLines = False
dark_bg(ws2, 60, 12)

write_title(ws2, 1, 1, "MONTHLY REVENUE REPORT — 2011 to 2015", 8)

hdrs = ["Year","Month","Year-Month",
        "Transactions","Total Revenue","Avg Revenue","Unique Customers"]
write_headers(ws2, 3, 1, hdrs)

for i, (_, r) in enumerate(monthly.iterrows(), 4):
    bg = CARD if i % 2 == 0 else HEADER
    write_row(ws2, i, 1, [
        r['year'], r['month_name'], r['year_month'],
        r['transactions'], r['total_revenue'],
        r['avg_revenue'],  r['unique_customers']
    ], bg)

# Line chart
chart = LineChart()
chart.title  = "Monthly Revenue Trend"
chart.style  = 10
chart.height = 14
chart.width  = 24
data = Reference(ws2, min_col=5, max_col=5,
                 min_row=3, max_row=3+len(monthly))
cats = Reference(ws2, min_col=3,
                 min_row=4, max_row=3+len(monthly))
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.series[0].graphicalProperties.line.solidFill = ORANGE
chart.series[0].graphicalProperties.line.width     = 25000
ws2.add_chart(chart, "I3")

set_widths(ws2, [8,12,14,14,16,14,18])
print("  ✔ Monthly Revenue sheet done")

# ══════════════════════════════════════════════
# SHEET 3 — YEARLY REVENUE
# ══════════════════════════════════════════════
ws3 = wb.create_sheet("Yearly Revenue")
ws3.sheet_view.showGridLines = False
dark_bg(ws3, 30, 12)

write_title(ws3, 1, 1, "YEARLY REVENUE REPORT — 2011 to 2015", 7)

hdrs3 = ["Year","Transactions","Total Revenue",
         "Avg Revenue","Unique Customers"]
write_headers(ws3, 3, 1, hdrs3)

for i, (_, r) in enumerate(yearly.iterrows(), 4):
    bg = CARD if i % 2 == 0 else HEADER
    write_row(ws3, i, 1, [
        r['year'], r['transactions'],
        r['total_revenue'], r['avg_revenue'],
        r['unique_customers']
    ], bg)

chart3 = BarChart()
chart3.title  = "Revenue by Year"
chart3.type   = "col"
chart3.style  = 10
chart3.height = 14
chart3.width  = 20
data3 = Reference(ws3, min_col=3, max_col=3,
                  min_row=3, max_row=3+len(yearly))
cats3 = Reference(ws3, min_col=1,
                  min_row=4, max_row=3+len(yearly))
chart3.add_data(data3, titles_from_data=True)
chart3.set_categories(cats3)
ws3.add_chart(chart3, "G3")

set_widths(ws3, [8,14,16,14,18])
print("  ✔ Yearly Revenue sheet done")

# ══════════════════════════════════════════════
# SHEET 4 — CUSTOMER SEGMENTS
# ══════════════════════════════════════════════
ws4 = wb.create_sheet("Customer Segments")
ws4.sheet_view.showGridLines = False
dark_bg(ws4, 20, 12)

write_title(ws4, 1, 1, "CUSTOMER SEGMENT ANALYSIS", 6)

hdrs4 = ["Segment","Customers","Transactions",
         "Total Revenue","Avg Revenue"]
write_headers(ws4, 3, 1, hdrs4)

for i, (_, r) in enumerate(segments.iterrows(), 4):
    bg = CARD if i % 2 == 0 else HEADER
    write_row(ws4, i, 1, [
        r['customer_segment'], r['customers'],
        r['transactions'], r['total_revenue'],
        r['avg_revenue']
    ], bg)

chart4 = BarChart()
chart4.title  = "Revenue by Segment"
chart4.type   = "bar"
chart4.style  = 10
chart4.height = 10
chart4.width  = 18
data4 = Reference(ws4, min_col=4, max_col=4,
                  min_row=3, max_row=3+len(segments))
cats4 = Reference(ws4, min_col=1,
                  min_row=4, max_row=3+len(segments))
chart4.add_data(data4, titles_from_data=True)
chart4.set_categories(cats4)
ws4.add_chart(chart4, "G3")

set_widths(ws4, [14,12,14,16,14])
print("  ✔ Customer Segments sheet done")

# ══════════════════════════════════════════════
# SHEET 5 — RFM ANALYSIS
# ══════════════════════════════════════════════
ws5 = wb.create_sheet("RFM Analysis")
ws5.sheet_view.showGridLines = False
dark_bg(ws5, 20, 10)

write_title(ws5, 1, 1, "RFM CUSTOMER SEGMENTATION", 7)

hdrs5 = ["Segment","Customers","Avg Recency (days)",
         "Avg Frequency","Avg Monetary","Total Revenue"]
write_headers(ws5, 3, 1, hdrs5)

for i, (_, r) in enumerate(rfm_summary.iterrows(), 4):
    bg = CARD if i % 2 == 0 else HEADER
    write_row(ws5, i, 1, [
        r['segment'],      r['customers'],
        r['avg_recency'],  r['avg_frequency'],
        r['avg_monetary'], r['total_revenue']
    ], bg)

set_widths(ws5, [18,12,20,15,16,16])
print("  ✔ RFM Analysis sheet done")

# ══════════════════════════════════════════════
# SHEET 6 — DAY OF WEEK
# ══════════════════════════════════════════════
ws6 = wb.create_sheet("Day of Week")
ws6.sheet_view.showGridLines = False
dark_bg(ws6, 15, 8)

write_title(ws6, 1, 1, "REVENUE BY DAY OF WEEK", 5)

hdrs6 = ["Day","Transactions","Total Revenue","Avg Revenue"]
write_headers(ws6, 3, 1, hdrs6)

for i, (_, r) in enumerate(dow.iterrows(), 4):
    bg = CARD if i % 2 == 0 else HEADER
    write_row(ws6, i, 1, [
        r['day_name'], r['transactions'],
        r['total_revenue'], r['avg_revenue']
    ], bg)

chart6 = BarChart()
chart6.title  = "Revenue by Day of Week"
chart6.type   = "col"
chart6.style  = 10
chart6.height = 12
chart6.width  = 18
data6 = Reference(ws6, min_col=3, max_col=3,
                  min_row=3, max_row=3+len(dow))
cats6 = Reference(ws6, min_col=1,
                  min_row=4, max_row=3+len(dow))
chart6.add_data(data6, titles_from_data=True)
chart6.set_categories(cats6)
ws6.add_chart(chart6, "F3")

set_widths(ws6, [14,14,16,14])
print("  ✔ Day of Week sheet done")

# ══════════════════════════════════════════════
# SHEET 7 — RESPONSE ANALYSIS
# ══════════════════════════════════════════════
ws7 = wb.create_sheet("Response Analysis")
ws7.sheet_view.showGridLines = False
dark_bg(ws7, 15, 8)

write_title(ws7, 1, 1, "CAMPAIGN RESPONSE ANALYSIS", 5)

hdrs7 = ["Responded","Customers","Transactions",
         "Total Revenue","Avg Revenue"]
write_headers(ws7, 3, 1, hdrs7)

for i, (_, r) in enumerate(resp_summary.iterrows(), 4):
    bg = CARD if i % 2 == 0 else HEADER
    write_row(ws7, i, 1, [
        r['responded'],    r['customers'],
        r['transactions'], r['total_revenue'],
        r['avg_revenue']
    ], bg)

set_widths(ws7, [12,12,14,16,14])
print("  ✔ Response Analysis sheet done")

# ══════════════════════════════════════════════
# SHEET 8 — TOP CUSTOMERS
# ══════════════════════════════════════════════
ws8 = wb.create_sheet("Top Customers")
ws8.sheet_view.showGridLines = False
dark_bg(ws8, 30, 10)

write_title(ws8, 1, 1, "TOP 20 CUSTOMERS BY REVENUE", 7)

hdrs8 = ["Rank","Customer ID","Transactions",
         "Total Spent","Avg Spent",
         "First Purchase","Last Purchase"]
write_headers(ws8, 3, 1, hdrs8)

for i, (_, r) in enumerate(top_custs.iterrows(), 4):
    bg = CARD if i % 2 == 0 else HEADER
    write_row(ws8, i, 1, [
        i-3, r['customer_id'],
        r['transactions'],   r['total_spent'],
        r['avg_spent'],
        str(r['first_purchase'])[:10],
        str(r['last_purchase'])[:10]
    ], bg)

set_widths(ws8, [6,14,14,14,12,16,16])
print("  ✔ Top Customers sheet done")

# ══════════════════════════════════════════════
# SHEET 9 — QUARTERLY
# ══════════════════════════════════════════════
ws9 = wb.create_sheet("Quarterly")
ws9.sheet_view.showGridLines = False
dark_bg(ws9, 30, 10)

write_title(ws9, 1, 1, "QUARTERLY REVENUE BREAKDOWN", 6)

hdrs9 = ["Quarter","Year","Q",
         "Transactions","Total Revenue","Avg Revenue"]
write_headers(ws9, 3, 1, hdrs9)

for i, (_, r) in enumerate(quarterly.iterrows(), 4):
    bg = CARD if i % 2 == 0 else HEADER
    write_row(ws9, i, 1, [
        r['label'],        r['year'],
        r['quarter'],      r['transactions'],
        r['total_revenue'],r['avg_revenue']
    ], bg)

chart9 = BarChart()
chart9.title  = "Quarterly Revenue"
chart9.type   = "col"
chart9.style  = 10
chart9.height = 12
chart9.width  = 18
data9 = Reference(ws9, min_col=5, max_col=5,
                  min_row=3, max_row=3+len(quarterly))
cats9 = Reference(ws9, min_col=1,
                  min_row=4, max_row=3+len(quarterly))
chart9.add_data(data9, titles_from_data=True)
chart9.set_categories(cats9)
ws9.add_chart(chart9, "H3")

set_widths(ws9, [14,8,6,14,16,14])
print("  ✔ Quarterly sheet done")


wb.save(OUT_FILE)

print(f"   File saved: {OUT_FILE}")
print(f"\n   Sheets created:")
for s in wb.sheetnames:
    print(f"   ✔ {s}")